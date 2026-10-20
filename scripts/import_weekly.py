#!/usr/bin/env python3
"""Convertit des exports de programme hebdomadaire (xlsx : Routing, Num Vol, Owner, semaines…)
en data/weekly.json. À lancer en local : les exports bruts ne doivent pas être publiés.

    python3 scripts/import_weekly.py export_DL.xlsx export_KL.xlsx export_TO.xlsx export_AF.xlsx

Si l'export contient « Dep LT » / « Arr LT » (heures locales), la durée réelle est calculée avec
le fuseau horaire de chaque aéroport. Sinon la durée reste vide et build.py l'estime
(vols observés OpenSky, ou modèle distance -> temps de bloc).
Chaque lancement remplace complètement data/weekly.json : passe tous les exports à jour en une fois.
"""
import csv
import datetime as dt
import json
import re
import sys
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TZ_URL = 'https://raw.githubusercontent.com/jpatokal/openflights/master/data/airports.dat'

# Codes avion IATA (exports) -> codes OACI (Simbrief)
AC_IATA = {
    '332': 'A332', '333': 'A333', '338': 'A338', '339': 'A339', '359': 'A359', '351': 'A35K',
    '319': 'A319', '320': 'A320', '321': 'A321', '318': 'A318', '32N': 'A20N', '32Q': 'A21N',
    '223': 'BCS3', '221': 'BCS1', '772': 'B772', '77W': 'B77W', '77L': 'B77L', '788': 'B788',
    '789': 'B789', '781': 'B78X', '73H': 'B738', '73W': 'B737', '73J': 'B739', '7M8': 'B38M',
    '7M9': 'B39M', '752': 'B752', '753': 'B753', '763': 'B763', '764': 'B764', '717': 'B712',
    'E70': 'E170', 'E7W': 'E175', 'E75': 'E175', 'E90': 'E190', 'E95': 'E195', '290': 'E290',
    '295': 'E295', 'AT7': 'AT76', 'CR9': 'CRJ9',
}
WEEK = re.compile(r'^S\d+\s+(\d{2})/(\d{2})/(\d{4})')


def load_tz():
    path = ROOT / 'oa_airports.dat'
    if not path.exists():
        urllib.request.urlretrieve(TZ_URL, path)
    tz = {}
    with open(path, encoding='utf-8') as f:
        for r in csv.reader(f):
            if len(r) > 11 and len(r[4]) == 3 and r[11] not in ('\\N', ''):
                tz[r[4]] = r[11]
    # aéroports récents absents de la base OpenFlights
    tz.update({'IST': 'Europe/Istanbul', 'BER': 'Europe/Berlin', 'HYD': 'Asia/Kolkata'})
    return tz


def parse_time(s):
    """'07:45 (+1)' -> (465, 1)"""
    m = re.match(r'\s*(\d{1,2}):(\d{2})\s*(?:\(([+-]\d)\))?', str(s or ''))
    if not m:
        return None
    return int(m.group(1)) * 60 + int(m.group(2)), int(m.group(3) or 0)


def first_operating_date(row, weeks, today):
    """Premier jour d'opération à partir d'aujourd'hui (sert à choisir heure d'été/hiver)."""
    best, total = None, 0
    for idx, start in weeks:
        mask = str(row[idx] or '')
        for ch in mask:
            if ch.isdigit():
                total += 1
                day = start + dt.timedelta(days=int(ch) - 1)
                if day >= today and (best is None or day < best):
                    best = day
    return best, total


def main(paths):
    tz = load_tz()
    today = dt.date.today()
    agg = defaultdict(lambda: {'dur': [], 'dep': [], 'ac': Counter(), 'fn': Counter(), 'n': 0})
    no_tz = Counter()
    for path in paths:
        wb = load_workbook(path, read_only=True)
        rows = wb.active.iter_rows(values_only=True)
        head = [str(h or '').replace('\n', ' ').strip() for h in next(rows)]
        col = {h: i for i, h in enumerate(head)}
        weeks = []
        for i, h in enumerate(head):
            m = WEEK.match(h)
            if m:
                weeks.append((i, dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))))
        i_rt, i_fn, i_ow = col['Routing'], col['Num Vol'], col['Owner']
        i_ac = next((i for h, i in col.items() if h.replace(' ', '').lower() == 'typeavion'), None)
        i_dep, i_arr = col.get('Dep LT'), col.get('Arr LT')
        for row in rows:
            if not row or not row[i_rt] or not row[i_ow]:
                continue
            legs = str(row[i_rt]).strip().split('-')
            if len(legs) < 2 or not all(len(x) == 3 and x.isalpha() for x in legs):
                continue
            day, n_ops = first_operating_date(row, weeks, today)
            if not n_ops:
                continue
            op, fn = str(row[i_ow]).strip(), str(row[i_fn]).strip()
            ac = AC_IATA.get(str(row[i_ac]).strip(), str(row[i_ac]).strip()) if i_ac is not None and row[i_ac] else ''
            for a, b in zip(legs, legs[1:]):
                r = agg[(op, a, b)]
                r['n'] += n_ops
                r['fn'][fn] += n_ops
                if ac:
                    r['ac'][ac] += n_ops
            # Durée réelle seulement pour un vol direct avec heures locales
            if len(legs) == 2 and i_dep is not None and day:
                td, ta = parse_time(row[i_dep]), parse_time(row[i_arr])
                a, b = legs
                if td and ta and a in tz and b in tz:
                    d0 = dt.datetime.combine(day, dt.time()) + dt.timedelta(minutes=td[0])
                    d1 = dt.datetime.combine(day + dt.timedelta(days=ta[1]), dt.time()) + dt.timedelta(minutes=ta[0])
                    u0 = d0.replace(tzinfo=ZoneInfo(tz[a])).astimezone(dt.timezone.utc)
                    u1 = d1.replace(tzinfo=ZoneInfo(tz[b])).astimezone(dt.timezone.utc)
                    dur = round((u1 - u0).total_seconds() / 60)
                    if 15 <= dur <= 20 * 60:
                        agg[(op, a, b)]['dur'] += [dur] * n_ops
                        agg[(op, a, b)]['dep'] += [u0.hour * 60 + u0.minute] * n_ops
                elif td:
                    no_tz.update([x for x in legs if x not in tz])

    out = []
    for (op, a, b), r in sorted(agg.items()):
        fn = r['fn'].most_common(1)[0][0]
        prefix, digits = fn[:2], fn[2:].lstrip('0') or '0'
        dep = Counter(r['dep']).most_common(1)[0][0] if r['dep'] else None
        out.append({'op': op, 'fn': prefix + digits, 'dep': a, 'arr': b,
                    'dur': round(sum(r['dur']) / len(r['dur'])) if r['dur'] else None,
                    'hdep': f'{dep // 60:02d}:{dep % 60:02d}' if dep is not None else '',
                    'ac': r['ac'].most_common(1)[0][0] if r['ac'] else '', 'n': r['n']})
    dest = ROOT / 'data' / 'weekly.json'
    dest.write_text(json.dumps(out, ensure_ascii=False, separators=(',', ':')))
    print(f'{len(out)} routes -> {dest}')
    for op, c in Counter(r['op'] for r in out).most_common():
        timed = sum(1 for r in out if r['op'] == op and r['dur'])
        print(f'  {op} : {c} routes, dont {timed} avec horaires')
    if no_tz:
        print('Aéroports sans fuseau horaire :', dict(no_tz))


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1:])
